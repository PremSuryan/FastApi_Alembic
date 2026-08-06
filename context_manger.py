file_path = r"E:\Fast Api Alembic\name.json"

with open(file_path, "r") as f:
    import json
    data = json.load(f)

print(data)


#Excel 

excel_path = r"E:\Fast Api Alembic\name.xlsx"

from openpyxl import load_workbook

workbook = load_workbook(excel_path)

try:
    sheet = workbook.active

    for rows in sheet.iter_rows(values_only=True):
        for cell in rows:
            print(cell.value)

finally:
    workbook.close()